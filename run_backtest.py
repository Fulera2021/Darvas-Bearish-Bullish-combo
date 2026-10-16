"""CLI Runner for the 5-Tranche Weekly Breakout Backtesting Engine.
Executes strategy simulations for single or batch stocks and exports Excel reports.
"""

import argparse
import logging
import os
import sys
from datetime import datetime
import pandas as pd

from config import BacktestConfig
from data_loader import fetch_daily_data, fetch_dividend_data, build_weekly_data, add_weekly_signals
from backtest_engine import run_backtest, BacktestResult
from excel_reporter import export_backtest_to_excel

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("WeeklyBreakoutRunner")

def run_single_stock(symbol: str, config: BacktestConfig, output_dir: str) -> BacktestResult:
    """Runs the strategy backtest on a single stock and generates an Excel workbook."""
    logger.info(f"=== Starting Backtest for {symbol} ({config.start_date} to {config.end_date}) ===")
    
    # 1. Fetch Daily Data with 365-day warmup for rolling calculations and cold start
    warmup_start_dt = (pd.to_datetime(config.start_date) - pd.Timedelta(days=365)).strftime("%Y-%m-%d")
    df_daily = fetch_daily_data(symbol, warmup_start_dt, config.end_date)
    logger.info(f"Loaded {len(df_daily)} daily bars (warmup from {warmup_start_dt}).")

    # 2. Resample to Weekly and Compute Breakout Signals
    df_weekly = build_weekly_data(df_daily)
    df_weekly = add_weekly_signals(
        df_weekly,
        entry_buffer_pct=config.entry_buffer_pct,
        use_trend_filter=config.use_trend_filter,
        trend_ema_period=config.trend_ema_period,
        darvas_box_weeks=config.darvas_box_weeks
    )
    signals_count = int(df_weekly["breakout_signal"].sum())
    logger.info(f"Resampled to {len(df_weekly)} weekly bars. Detected {signals_count} breakout signals.")

    # Fetch dividends if available
    df_div = fetch_dividend_data(symbol)

    # 3. Run Event-Driven Simulator
    result = run_backtest(df_daily, df_weekly, config, symbol=symbol, df_dividends=df_div)

    # 4. Print Summary Console Output (Plain numbers, no rupee symbols)
    print("\n" + "=" * 65)
    print(f"  BACKTEST RESULTS: {symbol}")
    print("=" * 65)
    print(f"  Initial Portfolio Capital    : {result.initial_capital:>15,.2f}")
    print(f"  Capital Per Tranche (1/5)    : {config.tranche_capital:>15,.2f}")
    if config.enable_magic_profit:
        print(f"  Trade Realized Profit        : {result.realized_trade_profit:>15,.2f}")
        print(f"  Magic Profit (+500 rule)     : {result.total_magic_profit:>15,.2f}")
        print(f"  Magic Profit Events Count    : {len(result.magic_events):>15}")
    print(f"  Total Realized Profit        : {result.total_realized_profit:>15,.2f}")
    print(f"  Unrealized P&L (Open)        : {result.unrealized_pnl:>15,.2f}")
    print(f"  Total Portfolio Return       : {result.total_portfolio_return_pct:>14.2f}%")
    print(f"  XIRR % (Deployed Capital)    : {result.deployed_xirr_pct:>14.2f}%")
    print(f"  Max Drawdown                 : {result.max_drawdown_pct:>14.2f}%")
    print(f"  Max Concurrent Tranches Used : {result.max_active_tranches:>15} / {config.num_parts}")
    print(f"  Total Trades Taken           : {result.total_trades:>15}")
    print(f"  Completed (Closed) Trades    : {result.completed_trades:>15}")
    print(f"  Open Trades Currently Active : {result.open_trades:>15}")
    print(f"  Average Holding Duration     : {result.avg_holding_days:>14.1f} calendar days")
    print("=" * 65 + "\n")

    # 5. Export Excel Workbook
    clean_sym = symbol.replace(".NS", "").replace("^", "")
    os.makedirs(output_dir, exist_ok=True)
    if config.enable_bullish_combo:
        mode_tag = "Combo"
    elif config.use_20w_low_tracking:
        mode_tag = "20W_Tracking"
    elif config.enable_magic_profit:
        mode_tag = "Magic"
    elif config.buy_only_below_last_tranche:
        mode_tag = "BelowLast"
    elif config.use_weekly_trailing_exit:
        mode_tag = "Trailing"
    else:
        mode_tag = "Base"

    out_file = os.path.join(output_dir, f"Weekly_Breakout_{mode_tag}_{clean_sym}.xlsx")
    export_backtest_to_excel(result, config, out_file)
    logger.info(f"Excel report generated: {out_file}")
    
    return result

def main():
    parser = argparse.ArgumentParser(description="5-Tranche Weekly Breakout Swing Backtester")
    parser.add_argument("--symbol", type=str, default="RELIANCE.NS", help="NSE stock ticker (e.g. RELIANCE.NS, TCS.NS)")
    parser.add_argument("--start", type=str, default="2021-01-01", help="Backtest start date (YYYY-MM-DD)")
    parser.add_argument("--end", type=str, default=datetime.today().strftime("%Y-%m-%d"), help="Backtest end date (YYYY-MM-DD)")
    parser.add_argument("--capital", type=float, default=100000.0, help="Total capital (default: 100000)")
    parser.add_argument("--parts", type=int, default=5, help="Number of capital parts (default: 5)")
    parser.add_argument("--target", type=float, default=0.05, help="Profit target per tranche (default: 0.05 for 5%)")
    parser.add_argument("--buffer", type=float, default=0.005, help="Next week entry price ceiling buffer (default: 0.005 for 0.5%)")
    parser.add_argument("--combo", action="store_true", help="Enable Darvas Bearish dip buying + Bullish momentum combo")
    parser.add_argument("--bullish-min-return", type=float, default=15.0, help="Min return % p.a. for bullish tranche exit (default: 15.0)")
    parser.add_argument("--trailing", action="store_true", help="Use weekly trailing low exit (sell at Friday Close if Close < Prev Week Low and in profit)")
    parser.add_argument("--buy-below-last", action="store_true", help="Only buy if signal weekly close is below the last bought tranche's buy price")
    parser.add_argument("--magic", action="store_true", help="Enable magic profit (+500 on 20DL breakdown while in loss, cooldown 30 days)")
    parser.add_argument("--tracking-20w", action="store_true", help="Enable 20-week low breakdown tracking to weekly breakout entry until 20-week high")
    parser.add_argument("--batch", action="store_true", help="Run backtest across all default watchlist symbols")
    parser.add_argument("--output-dir", type=str, default=None, help="Directory to save generated Excel workbooks")
    
    args = parser.parse_args()

    enable_bullish = False
    if args.combo:
        default_subfolder = "Bearish + Bullish Combo"
        mode_tag = "Combo"
        buy_below_last = False
        enable_magic = False
        use_trailing = False
        use_20w_tracking = True
        enable_bullish = True
    elif args.tracking_20w:
        default_subfolder = "20W Low Tracking"
        mode_tag = "20W_Tracking"
        buy_below_last = False
        enable_magic = False
        use_trailing = False
        use_20w_tracking = True
    elif args.magic:
        default_subfolder = "Base + buy low + magic"
        mode_tag = "Magic"
        buy_below_last = True
        enable_magic = True
        use_trailing = False
        use_20w_tracking = False
    elif args.buy_below_last:
        default_subfolder = "Base + buying below last tranche"
        mode_tag = "BelowLast"
        buy_below_last = True
        enable_magic = False
        use_trailing = False
        use_20w_tracking = False
    elif args.trailing:
        default_subfolder = "trailing"
        mode_tag = "Trailing"
        buy_below_last = False
        enable_magic = False
        use_trailing = True
        use_20w_tracking = False
    else:
        default_subfolder = "base"
        mode_tag = "Base"
        buy_below_last = False
        enable_magic = False
        use_trailing = False
        use_20w_tracking = False

    current_dir = os.path.dirname(os.path.abspath(__file__))
    output_dir = args.output_dir or os.path.join(current_dir, "output", default_subfolder)

    config = BacktestConfig(
        total_capital=args.capital,
        num_parts=args.parts,
        target_pct=args.target,
        entry_buffer_pct=args.buffer,
        use_weekly_trailing_exit=use_trailing,
        buy_only_below_last_tranche=buy_below_last,
        enable_magic_profit=enable_magic,
        use_20w_low_tracking=use_20w_tracking,
        enable_bullish_combo=enable_bullish,
        bullish_min_return_pct=args.bullish_min_return,
        start_date=args.start,
        end_date=args.end
    )

    if args.batch:
        summary_rows = []
        for sym in config.symbols:
            try:
                res = run_single_stock(sym, config, output_dir)
                if enable_magic:
                    summary_rows.append({
                        "Symbol": sym.replace(".NS", ""),
                        "Total Capital": res.initial_capital,
                        "Total Realized Profit": res.total_realized_profit,
                        "Trade Realized Profit": res.realized_trade_profit,
                        "Magic Profit": res.total_magic_profit,
                        "Unrealized P&L": res.unrealized_pnl,
                        "Total Return %": res.total_portfolio_return_pct / 100.0,
                        "XIRR %": res.deployed_xirr_pct / 100.0,
                        "Max Drawdown %": res.max_drawdown_pct / 100.0,
                        "Max Tranches Used": res.max_active_tranches,
                        "Total Trades": res.total_trades,
                        "Completed Trades": res.completed_trades,
                        "Open Trades": res.open_trades,
                        "Magic Events": len(res.magic_events),
                        "Avg Holding Days": res.avg_holding_days
                    })
                else:
                    summary_rows.append({
                        "Symbol": sym.replace(".NS", ""),
                        "Total Capital": res.initial_capital,
                        "Realized Profit": res.total_realized_profit,
                        "Unrealized P&L": res.unrealized_pnl,
                        "Total Return %": res.total_portfolio_return_pct / 100.0,
                        "XIRR %": res.deployed_xirr_pct / 100.0,
                        "Max Drawdown %": res.max_drawdown_pct / 100.0,
                        "Max Tranches Used": res.max_active_tranches,
                        "Total Trades": res.total_trades,
                        "Completed Trades": res.completed_trades,
                        "Open Trades": res.open_trades,
                        "Avg Holding Days": res.avg_holding_days
                    })
            except Exception as e:
                logger.error(f"Error processing {sym}: {e}", exc_info=True)
                
        if summary_rows:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            summary_path = os.path.join(output_dir, f"Watchlist_Consolidated_Summary_{mode_tag}.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Consolidated Summary"
            ws.views.sheetView[0].showGridLines = True

            font_title = Font(name="Segoe UI", size=14, bold=True, color="1F497D")
            font_sub = Font(name="Segoe UI", size=10, italic=True, color="595959")
            font_hdr = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            font_body = Font(name="Segoe UI", size=9)
            font_bold = Font(name="Segoe UI", size=9, bold=True)
            fill_hdr = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            fill_zebra = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")
            thin_side = Side(border_style="thin", color="D9D9D9")
            border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

            ws["B2"] = f"5-TRANCHE STRATEGY - CONSOLIDATED WATCHLIST SUMMARY ({mode_tag.upper()})"
            ws["B2"].font = font_title
            ws["B3"] = f"Watchlist: Nifty 50 Top 5 | Test Period: {config.start_date} to {config.end_date} | Base Capital: 100,000 per Stock"
            ws["B3"].font = font_sub

            headers = list(summary_rows[0].keys())
            tbl_start = 5
            for c_idx, h in enumerate(headers, start=2):
                cell = ws.cell(row=tbl_start, column=c_idx, value=h)
                cell.font = font_hdr
                cell.fill = fill_hdr
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border_cell

            for r_idx, row_dict in enumerate(summary_rows, start=tbl_start + 1):
                is_zebra = (r_idx % 2 == 0)
                row_fill = fill_zebra if is_zebra else None
                for c_idx, h in enumerate(headers, start=2):
                    val = row_dict[h]
                    cell = ws.cell(row=r_idx, column=c_idx, value=val)
                    cell.font = font_body
                    cell.border = border_cell
                    if row_fill:
                        cell.fill = row_fill

                    if h == "Symbol":
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                        cell.font = font_bold
                    elif "%" in h:
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = "0.00%"
                    elif "Capital" in h or "Profit" in h or "P&L" in h:
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = "#,##0.00"
                    elif "Days" in h:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.number_format = "0.0"
                    else:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.number_format = "#,##0"

            ws.freeze_panes = f"B{tbl_start + 1}"
            for col in ws.columns:
                col_letter = get_column_letter(col[0].column)
                if col[0].column == 1:
                    ws.column_dimensions[col_letter].width = 3
                    continue
                max_len = max((len(str(c.value or "")) for c in col if c.row >= tbl_start), default=12)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

            try:
                wb.save(summary_path)
            except PermissionError:
                base, ext = os.path.splitext(summary_path)
                summary_path = f"{base}_new{ext}"
                wb.save(summary_path)

            logger.info(f"Consolidated Watchlist summary saved to: {summary_path}")
    else:
        run_single_stock(args.symbol, config, output_dir)

if __name__ == "__main__":
    main()
