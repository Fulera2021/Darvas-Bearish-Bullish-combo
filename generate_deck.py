"""
Script to generate a professional, modern 16:9 PowerPoint presentation
explaining the Darvas Bearish + Bullish Combo Systematic Backtesting Strategy
with clear visual infographics, charts, and live INFY (Infosys) case study.
"""

import os
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_filename="Darvas_Strategy_Rulebook_Deck_v2.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank layout

    # Color Palette Definitions
    BG_DARK = RGBColor(15, 23, 42)        # Slate 900
    BG_LIGHT = RGBColor(248, 250, 252)    # Slate 50
    CARD_BG = RGBColor(255, 255, 255)     # White
    CARD_BORDER = RGBColor(226, 232, 240) # Slate 200
    
    TEXT_DARK = RGBColor(15, 23, 42)      # Slate 900
    TEXT_MUTED = RGBColor(100, 116, 139)  # Slate 500
    TEXT_LIGHT = RGBColor(241, 245, 249)  # Slate 100
    
    BRAND_BLUE = RGBColor(37, 99, 235)    # Blue 600
    BRAND_DARK_BLUE = RGBColor(30, 58, 138) # Blue 900
    EMERALD_GREEN = RGBColor(16, 185, 129)# Green 500
    EMERALD_BG = RGBColor(236, 253, 245)  # Green 50
    AMBER_ORANGE = RGBColor(217, 119, 6)  # Amber 600
    AMBER_BG = RGBColor(254, 243, 199)    # Amber 100
    PURPLE_INDIGO = RGBColor(99, 102, 241)# Indigo 500

    def add_header(slide, title_text, category_text="OPERATIONAL RULEBOOK"):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = BRAND_BLUE
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK
        
        # Subtle horizontal divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.color.rgb = CARD_BORDER

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Theme with Badges & Emojis)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = BG_DARK
    bg1.line.color.rgb = BG_DARK

    # Top Pill Badge
    badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(5.2), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(30, 41, 59)
    badge.line.color.rgb = RGBColor(51, 65, 85)
    tb = badge.text_frame
    tb.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tb.paragraphs[0]
    p.text = "★ SYSTEMATIC TRADING & BACKTESTING ENGINE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248) # Cyan

    # Title & Subtitle
    tbox = slide1.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.3), Inches(2.8))
    tf = tbox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "The Darvas Bearish + Bullish Combo"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = tf.add_paragraph()
    p2.text = "Operational Rulebook & Behind-the-Scenes Backtesting Mechanics\n4-Tranche Capital Architecture (Rs. 40,000 Portfolio) | Live Case Study: INFY (Infosys Ltd)"
    p2.font.size = Pt(17)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.space_before = Pt(10)

    # 3 Summary Pill Boxes with Emojis
    pills = [
        ("📡 20-WEEK RADAR", "Rolling 20-Week Low & High trigger automated state detection without emotion.", BRAND_BLUE),
        ("💼 4 TRANCHE SLOTS", "Rs. 40,000 capital split into 4 parts (1 Bullish Trend + 3 Bearish Dip Harvesters).", EMERALD_GREEN),
        ("🇮🇳 LIVE INFY PROOF", "100% win rate on 33 closed trades, +Rs. 40,724 cash profit (+101.8% gain) from 2020-2026!", AMBER_ORANGE)
    ]
    for i, (title, desc, color) in enumerate(pills):
        card = add_card(slide1, Inches(1.0 + i * 3.9), Inches(5.0), Inches(3.6), Inches(1.8), RGBColor(30, 41, 59), RGBColor(51, 65, 85))
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.2)
        ctf.margin_top = Inches(0.2)
        cp0 = ctf.paragraphs[0]
        cp0.text = title
        cp0.font.size = Pt(13)
        cp0.font.bold = True
        cp0.font.color.rgb = color
        cp1 = ctf.add_paragraph()
        cp1.text = desc
        cp1.font.size = Pt(10.5)
        cp1.font.color.rgb = RGBColor(203, 213, 225)
        cp1.space_before = Pt(6)

    # =========================================================================
    # SLIDE 2: Core Philosophy & Dual Engine Concept
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "The Core Architecture: Two Complementary Engines in One System")

    # Left Column: The Market Problem (Cards)
    add_card(slide2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3))
    tbox = slide2.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.9))
    tf = tbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚠️ THE TRADITIONAL INVESTING DILEMMA"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = AMBER_ORANGE
    
    bullets = [
        ("Markets Chop 70% of the Time 📉", "Stocks rarely go straight up. Most of the year is spent in frustrating pullbacks, corrections, and sideways channels."),
        ("Traditional Buy-and-Hold Gets Trapped ⏳", "Investors sit on idle capital or suffer severe drawdowns waiting years for a stock to recover from cyclical dips."),
        ("Pure Momentum Strategies Get Whiplashed 🌪️", "Buying only at new highs often results in buying at market tops and getting panicked into selling at bottoms.")
    ]
    for h, b in bullets:
        ph = tf.add_paragraph()
        ph.text = h
        ph.font.size = Pt(13)
        ph.font.bold = True
        ph.font.color.rgb = TEXT_DARK
        ph.space_before = Pt(14)
        pb = tf.add_paragraph()
        pb.text = b
        pb.font.size = Pt(11)
        pb.font.color.rgb = TEXT_MUTED

    # Right Column: The Dual Engine Solution
    add_card(slide2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3))
    tbox2 = slide2.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.9))
    tf2 = tbox2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🛡️ OUR DUAL-ENGINE SOLUTION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE

    solutions = [
        ("ENGINE 1: Bearish Dip-Harvester (Slots 2 to 4) 🎯", "When a stock enters a correction, our radar activates. Each time the stock makes a weekly swing bounce, we buy a tranche and exit quickly at a strict +5% target. It turns market drops into cash generators!"),
        ("ENGINE 2: Bullish Trend-Rider (Slot 1) 🚀", "When a stock breaks above its 20-week high, we allocate 1 dedicated slot to ride the massive macro uptrend. We do NOT cap it at 5%; we let it compound until the entire multi-year bull run finishes!"),
        ("Dynamic Compounding & Recycling 🔄", "Every profit won and dividend paid is automatically reinvested back into future tranches, accelerating wealth exponentially without adding fresh capital.")
    ]
    for h, b in solutions:
        ph = tf2.add_paragraph()
        ph.text = h
        ph.font.size = Pt(13)
        ph.font.bold = True
        ph.font.color.rgb = TEXT_DARK
        ph.space_before = Pt(14)
        pb = tf2.add_paragraph()
        pb.text = b
        pb.font.size = Pt(11)
        pb.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 3: Capital Architecture & Tranche Sizing (Embed Graphic)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Capital Architecture: How Rs. 40,000 is Divided into 4 Tranches")

    # Embed generated tranche architecture infographic
    if os.path.exists("graphics/tranche_architecture.png"):
        slide3.shapes.add_picture("graphics/tranche_architecture.png", Inches(0.8), Inches(1.5), width=Inches(11.733))

    # Bottom Explanation Cards
    add_card(slide3, Inches(0.8), Inches(5.2), Inches(5.7), Inches(1.8))
    tb_b1 = slide3.shapes.add_textbox(Inches(1.0), Inches(5.3), Inches(5.3), Inches(1.6))
    tf = tb_b1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🧮 DISCRETE SHARE SIZING (NO MARGIN)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE
    p1 = tf.add_paragraph()
    p1.text = "Shares Bought = Floor(Tranche Capital / Buy Price)\n• Example: Tranche Capital = Rs. 10,000 and Stock is Rs. 746:\n  Shares = Floor(10,000 / 746) = 13 Shares (Invested: Rs. 9,698)\n• Unused fractional cash stays 100% safe in bank buffer earning interest!"
    p1.font.size = Pt(10)
    p1.font.color.rgb = TEXT_DARK
    p1.space_before = Pt(4)

    add_card(slide3, Inches(6.8), Inches(5.2), Inches(5.7), Inches(1.8))
    tb_b2 = slide3.shapes.add_textbox(Inches(7.0), Inches(5.3), Inches(5.3), Inches(1.6))
    tf = tb_b2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📈 DYNAMIC CAPITAL COMPOUNDING"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN
    p1 = tf.add_paragraph()
    p1.text = "Current Capital = Initial (Rs. 40k) + Realized Profits + Dividends\nTranche Capital = Current Capital / 4 Parts\n• As profits roll in, Slot size automatically grows to Rs. 12k, 15k, 18k+!\n• Each slot operates with bigger firepower without depositing any fresh money."
    p1.font.size = Pt(10)
    p1.font.color.rgb = TEXT_DARK
    p1.space_before = Pt(4)

    # =========================================================================
    # SLIDE 4: Phase 1 – When Tracking Starts (Embed Workflow Graphic)
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Phase 1: When Does Tracking Start? (The 20-Week Radar Mechanics)")

    # 3 Horizontal Step Cards
    steps = [
        ("🔭 The 20-Week Window", "Calculates two rolling benchmarks every Friday:\n• 20-Week High (20WH): Resistance Ceiling\n• 20-Week Low (20WL): Support Floor\nCovers 5 months (~institutional cycle).", BRAND_BLUE),
        ("🚨 Bearish Radar Activation", "When Friday Close < 20-Week Low:\n• Stock is confirmed in a deep correction.\n• Bearish Radar turns ON.\n• Begins hunting weekly dip-buying bounces.", AMBER_ORANGE),
        ("🚀 Bullish Breakout Shift", "When Friday Close > 20-Week High:\n• Bearish Radar immediately turns OFF.\n• Deep dip phase is over; bull trend starts.\n• Slot 1 Bullish Tranche deployed to ride it!", EMERALD_GREEN)
    ]

    for i, (title, desc, col) in enumerate(steps):
        card = add_card(slide4, Inches(0.8 + i * 3.95), Inches(1.6), Inches(3.8), Inches(1.9), CARD_BG, col)
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.2)
        ctf.margin_top = Inches(0.15)
        
        p = ctf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col
        
        p1 = ctf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(10)
        p1.font.color.rgb = TEXT_DARK
        p1.space_before = Pt(4)

    # Embed visual workflow diagram below
    if os.path.exists("graphics/radar_workflow.png"):
        slide4.shapes.add_picture("graphics/radar_workflow.png", Inches(0.8), Inches(3.7), width=Inches(11.733))

    # =========================================================================
    # SLIDE 5: Phase 2 – When Buying is Done & Execution Mechanics
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Phase 2: When is Buying Done? (Two-Step Execution Mechanics)")

    # Left Column: The 2-Step Trigger
    add_card(slide5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3))
    tb_l = slide5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.9))
    tf = tb_l.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🛒 TWO-STEP ORDER EXECUTION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE

    bullets_buy = [
        ("Step 1: Signal Detection (Friday 3:30 PM)", 
         "The backtest engine checks Friday Weekly Close:\n"
         "Does Weekly Close > Previous Week High?\n"
         "If YES and free slots exist, a Pending Buy Setup is queued.\n"
         "Entry Ceiling Price = Friday Close * 1.005 (+0.5% buffer)"),
        ("Step 2: Order Fill (Monday to Friday Daily Market)", 
         "During the next trading week, the engine checks every daily bar:\n"
         "Is Daily Low <= Entry Ceiling Price?\n"
         "If YES: Buy Tranche is FILLED!\n"
         "Fill Price = min(Daily Open, Entry Ceiling Price)\n"
         "Target Price = Buy Price * 1.05 (+5.0% profit)"),
        ("The Expiration Rule ⏳", 
         "If the stock gaps up violently and stays above the ceiling all week, the order expires with zero penalty. Cycle repeats cleanly next week.")
    ]
    for h, b in bullets_buy:
        ph = tf.add_paragraph()
        ph.text = f"• {h}"
        ph.font.size = Pt(12.5)
        ph.font.bold = True
        ph.font.color.rgb = TEXT_DARK
        ph.space_before = Pt(12)
        pb = tf.add_paragraph()
        pb.text = b
        pb.font.size = Pt(10.5)
        pb.font.color.rgb = TEXT_MUTED

    # Right Column: Safety Rules During Deployment
    add_card(slide5, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3))
    tb_r = slide5.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.9))
    tf2 = tb_r.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🛡️ FOUR CRITICAL RISK GUARDRAILS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN

    rules_buy = [
        ("1️⃣ Maximum 1 Tranche Per Week", "No matter how strong a single week looks, we never deploy more than 1 slot in the same week. This enforces temporal diversification across weeks."),
        ("2️⃣ Strict 4-Slot Capacity Ceiling", "If all 4 slots are occupied, the engine automatically ignores all new buy signals. We never borrow on margin or over-allocate."),
        ("3️⃣ Anti-Clustering Spacing", "Prevents buying two tranches at the exact same price. New tranches are spaced out across distinct price levels."),
        ("4️⃣ Limit Price Ceiling (No Slippage)", "Using a +0.5% limit ceiling ensures we never chase sudden emotional price spikes or enter at unfavorable prices.")
    ]
    for h, b in rules_buy:
        ph = tf2.add_paragraph()
        ph.text = h
        ph.font.size = Pt(12.5)
        ph.font.bold = True
        ph.font.color.rgb = TEXT_DARK
        ph.space_before = Pt(12)
        pb = tf2.add_paragraph()
        pb.text = b
        pb.font.size = Pt(10.5)
        pb.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 6: Phase 3 – When Are They Sold? (Exit Mechanics)
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Phase 3: When Are Tranches Sold? (The Two Distinct Exit Paths)")

    # Left Card: Bearish Tranche Exits
    add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3), AMBER_BG, AMBER_ORANGE)
    tb_b = slide6.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.9))
    tf = tb_b.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🎯 PATH A: BEARISH DIP TRANCHES (SLOTS 2 - 4)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = AMBER_ORANGE

    p1 = tf.add_paragraph()
    p1.text = "STRICT +5.0% PROFIT TARGET EXIT"
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_DARK
    p1.space_before = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "• Rule: Sold the instant Daily High >= Buy Price * 1.05\n" \
              "• Execution: Sells automatically at the exact Target Price.\n" \
              "• Holding Time: Averages 20 to 45 calendar days.\n" \
              "• Immediate Slot Recycling 🔄: The moment the 5% profit is booked, the capital plus profit returns to available cash, and that specific Slot becomes 100% free to catch the next signal!\n" \
              "• No Stop Loss: We do not panic-sell dips. In quality large-cap stocks like Infosys, pullbacks are temporary. Each tranche holds until its +5% profit is realized."
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_DARK
    p2.space_before = Pt(10)

    # Right Card: Bullish Tranche Exits
    add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3), EMERALD_BG, EMERALD_GREEN)
    tb_g = slide6.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.9))
    tf2 = tb_g.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🚀 PATH B: BULLISH MOMENTUM TRANCHE (SLOT 1)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = EMERALD_GREEN

    p1 = tf2.add_paragraph()
    p1.text = "MACRO TREND COMPOUNDER (NO 5% CAP!)"
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_DARK
    p1.space_before = Pt(8)

    p2 = tf2.add_paragraph()
    p2.text = "• Philosophy: In a multi-year bull run, exiting at 5% leaves massive wealth on the table. Slot 1 is designed to ride the entire wave!\n" \
              "• Exit Condition 1: Weekly Close drops below the Rolling 20-Week Low (A full trend reversal confirmed).\n" \
              "• Exit Condition 2: AND the achieved return must be >= 15% per annum for the time held:\n" \
              "  Required Profit % = 15% * (Days Held / 365.25)\n" \
              "• In INFY: Slot 1 bought at Rs. 903 in July 2020 and held until April 2022 to exit at Rs. 1,585, booking a massive +75.5% gain (+Rs. 8,187 profit)!"
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_DARK
    p2.space_before = Pt(10)

    # =========================================================================
    # SLIDE 7: Complete Trade Lifecycle (Visual Process Flow)
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "The Complete Step-by-Step Lifecycle of a Trade")

    cycle_steps = [
        ("1️⃣ WEEKLY RADAR", "Friday 3:30 PM\n• Checks Close > 20W Low\n• Checks Close > Prev High\n• Verifies free slots exist", BRAND_BLUE),
        ("2️⃣ PENDING ORDER", "Weekend Setup\n• Locks reference price\n• Sets entry ceiling (+0.5%)\n• Sized to 1/4th capital", PURPLE_INDIGO),
        ("3️⃣ DAILY FILL", "Next Week Mon-Fri\n• Low <= Entry Ceiling\n• Fills at min(Open, Ceiling)\n• Target set at +5%", BRAND_DARK_BLUE),
        ("4️⃣ TARGET HIT", "Profit Realized! 🎯\n• High >= Target Price\n• +5% gain locked in cash\n• Slot instantly recycled 🔄", EMERALD_GREEN)
    ]

    for i, (title, desc, col) in enumerate(cycle_steps):
        card = add_card(slide7, Inches(0.8 + i * 2.98), Inches(1.8), Inches(2.8), Inches(2.9), CARD_BG, col)
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.2)
        ctf.margin_top = Inches(0.2)
        
        p = ctf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col
        
        p1 = ctf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(11)
        p1.font.color.rgb = TEXT_DARK
        p1.space_before = Pt(8)

    # Big Summary Card below
    add_card(slide7, Inches(0.8), Inches(5.0), Inches(11.7), Inches(1.9))
    tb_bot = slide7.shapes.add_textbox(Inches(1.0), Inches(5.1), Inches(11.3), Inches(1.7))
    tf = tb_bot.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🔄 THE RECYCLING FLYWHEEL: WHY CONTINUOUS ROTATION BEATS STATIC BUY & HOLD"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE

    p1 = tf.add_paragraph()
    p1.text = "In traditional investing, you buy once and wait for years. In our 4-tranche system:\n" \
              "• Slot 2 buys INFY at Rs. 652, sells at Rs. 685 (Profit: Rs. 489 in 34 days) -> Slot 2 is recycled.\n" \
              "• Slot 2 buys INFY at Rs. 715, sells at Rs. 751 (Profit: Rs. 500 in 57 days) -> Slot 2 is recycled again.\n" \
              "• Slot 2 buys INFY at Rs. 748, sells at Rs. 785 (Profit: Rs. 523 in 11 days) -> Slot 2 is recycled again.\n" \
              "The same Rs. 10,000 capital was deployed 3 times in 4 months, harvesting profits repeatedly!"
    p1.font.size = Pt(10.5)
    p1.font.color.rgb = TEXT_DARK
    p1.space_before = Pt(4)

    # =========================================================================
    # SLIDE 8: Live Case Study – INFY Performance Scorecard (Embed Equity Chart)
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Live Case Study: INFY (Infosys Ltd) Performance Scorecard (2020-2026)")

    # 4 Big Stat Cards Across Top
    kpis = [
        ("💰 TOTAL CAPITAL", "Rs. 40,000", "Initial base allocation", BRAND_DARK_BLUE),
        ("🏆 TOTAL REALIZED PROFIT", "Rs. 40,724", "+101.8% cash profit booked!", EMERALD_GREEN),
        ("🎯 WIN RATE (CLOSED)", "100.0%", "33 out of 33 trades won!", EMERALD_GREEN),
        ("🛡️ MAX DRAWDOWN", "-20.6%", "Navigated Covid & Tech crash", AMBER_ORANGE)
    ]

    for i, (title, val, sub, col) in enumerate(kpis):
        card = add_card(slide8, Inches(0.8 + i * 2.98), Inches(1.5), Inches(2.8), Inches(1.7), CARD_BG, col)
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.15)
        ctf.margin_top = Inches(0.12)
        
        p = ctf.paragraphs[0]
        p.text = title
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED
        
        p1 = ctf.add_paragraph()
        p1.text = val
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = col
        p1.space_before = Pt(2)
        
        p2 = ctf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(2)

    # Left Side: Key Operational Metrics Card
    add_card(slide8, Inches(0.8), Inches(3.35), Inches(4.8), Inches(3.8))
    tb_m = slide8.shapes.add_textbox(Inches(0.95), Inches(3.45), Inches(4.5), Inches(3.6))
    tf = tb_m.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "OPERATIONAL METRICS FOR INFY"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE

    metrics_text = [
        ("Total Trades Taken", "35 Trades (1 Bullish Trend + 34 Bearish Dips)"),
        ("Completed Trades", "33 Trades (All 33 hit targets cleanly)"),
        ("Active Open Positions", "2 Trades (Compounding currently)"),
        ("Deployed Capital XIRR %", "24.7% Annualized Return on deployed cash!"),
        ("Total Portfolio Return %", "92.3% Net Return across entire capital"),
        ("Max Active Slots Used", "4 / 4 Slots (Full utilization during dips)")
    ]
    for lbl, val in metrics_text:
        p_row = tf.add_paragraph()
        p_row.text = f"• {lbl}: {val}"
        p_row.font.size = Pt(10)
        p_row.font.color.rgb = TEXT_DARK
        p_row.space_before = Pt(5)

    # Right Side: Embed Equity Compounding Curve Chart
    if os.path.exists("graphics/infy_equity_curve.png"):
        slide8.shapes.add_picture("graphics/infy_equity_curve.png", Inches(5.8), Inches(3.35), width=Inches(6.733))

    # =========================================================================
    # SLIDE 9: INFY Live Chart & Real Trade Anatomy (Embed Price Chart)
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "INFY Chart & Trade Execution: Anatomy of Real Trades")

    # Embed INFY Price Chart across top
    if os.path.exists("graphics/infy_price_chart.png"):
        slide9.shapes.add_picture("graphics/infy_price_chart.png", Inches(0.8), Inches(1.5), width=Inches(11.733))

    # Bottom Insights & Examples Card
    add_card(slide9, Inches(0.8), Inches(5.35), Inches(11.733), Inches(1.8))
    tb_ins = slide9.shapes.add_textbox(Inches(1.0), Inches(5.45), Inches(11.3), Inches(1.6))
    tf = tb_ins.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "KEY INSIGHTS FROM REAL INFY TRADES (2020-2026)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = BRAND_BLUE

    p1 = tf.add_paragraph()
    p1.text = "1. 🚀 Master Bullish Ride (Trade #10): Entered 17-Jul-2020 at Rs. 903.15, rode the post-Covid bull market, and exited on 22-Apr-2022 at Rs. 1,585.45 (+75.5% profit / +Rs. 8,187.60 in 644 days)!\n" \
              "2. ⚡ High-Velocity Dip Rotations: Trade #5 exited in 6 days (+Rs. 493.50), Trade #9 in 5 days (+Rs. 508.17), Trade #25 in 6 days (+Rs. 705.06), Trade #32 in 6 days (+Rs. 845.52)!\n" \
              "3. 📈 Compounding at Work: Trade #1 in 2020 made Rs. 484.90, while Trade #34 in 2026 made Rs. 904.16 (+86% larger profit!) because each tranche dynamically grew from Rs. 9.7k to Rs. 18.0k!"
    p1.font.size = Pt(9.8)
    p1.font.color.rgb = TEXT_DARK
    p1.space_before = Pt(4)

    # =========================================================================
    # SLIDE 10: Executive Summary & Why This Works for Investors
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "Executive Summary: Why This Strategy Excels for Investors")

    cards = [
        ("🤖 100% Objective & Emotionless", 
         "• Zero forecasting, zero news watching, zero tips.\n• Everything is 100% mathematical based on weekly closing prices and strict limits.\n• Removes fear during crashes and eliminates greed during rallies.", 
         BRAND_BLUE),
        ("🛡️ Asymmetric Risk Protection", 
         "• Strict 4-tranche cap prevents dangerous portfolio concentration.\n• Max drawdown on INFY was contained at -20.6% across 6 years.\n• Idle capital sits safely earning risk-free interest buffer.", 
         EMERALD_GREEN),
        ("🇮🇳 Built for Quality Indian Equities", 
         "• Large-cap Nifty 50 stocks (like Infosys) have immense cyclical bounce reliability.\n• Dual engine monetizes frequent 5% swings while letting long-term trends run.\n• Achieves 24.7% XIRR with extraordinary peace of mind.", 
         AMBER_ORANGE)
    ]

    for i, (title, desc, col) in enumerate(cards):
        card = add_card(slide10, Inches(0.8 + i * 3.95), Inches(1.6), Inches(3.8), Inches(3.3), CARD_BG, col)
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.2)
        ctf.margin_top = Inches(0.2)
        
        p = ctf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = col
        
        p1 = ctf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(10.5)
        p1.font.color.rgb = TEXT_DARK
        p1.space_before = Pt(10)

    # Bottom Sign-Off Banner
    bot_banner = add_card(slide10, Inches(0.8), Inches(5.1), Inches(11.7), Inches(1.8), BG_DARK, BG_DARK)
    b_tf = bot_banner.text_frame
    b_tf.word_wrap = True
    b_tf.margin_left = b_tf.margin_right = Inches(0.3)
    b_tf.margin_top = Inches(0.2)

    bp0 = b_tf.paragraphs[0]
    bp0.text = "THE BOTTOM LINE FOR STAKEHOLDERS"
    bp0.font.size = Pt(12)
    bp0.font.bold = True
    bp0.font.color.rgb = RGBColor(56, 189, 248) # Cyan

    bp1 = b_tf.add_paragraph()
    bp1.text = "This systematic strategy solves the single biggest flaw in equity investing: emotional decision-making. By combining high-velocity 5% swing harvesting during corrections with a patient multi-year trend-riding tranche, it delivers compounding wealth (+101.8% cash profit on INFY) with strict risk management."
    bp1.font.size = Pt(11)
    bp1.font.color.rgb = RGBColor(241, 245, 249)
    bp1.space_before = Pt(6)

    # Save presentation with auto-fallback if file is locked in PowerPoint
    target_path = output_filename
    base, ext = os.path.splitext(output_filename)
    clean_base = base.split('_v')[0]
    saved = False
    for attempt in range(10):
        try:
            prs.save(target_path)
            print(f"Presentation saved successfully to: {target_path}")
            saved = True
            break
        except PermissionError:
            target_path = f"{clean_base}_v{attempt + 3}{ext}"
            
    if not saved:
        print("Warning: Could not save due to file locks.")

if __name__ == "__main__":
    create_deck()
