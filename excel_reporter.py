"""Excel report generator for the 5-Tranche Weekly Breakout Backtesting Engine.
Generates an executive KPI summary card dashboard and a granular trade journal
using openpyxl with clean styling and strictly plain numeric formatting (no currency symbols).
"""

import os
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import BacktestConfig
from backtest_engine import BacktestResult, Tranche

def export_backtest_to_excel(
    result: BacktestResult,
    config: BacktestConfig,
    output_filepath: str
) -> str:
    """
    Exports the backtest results to an Excel workbook (.xlsx).
    Summary of investment is placed at the top, followed by the complete trade journal.
    All numbers are formatted as plain numbers without any rupee symbol.
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{result.symbol[:25]} Backtest"
    ws.views.sheetView[0].showGridLines = True

    # ---------------------------------------------------------
    # Styles definition
    # ---------------------------------------------------------
    font_family = "Segoe UI"
    
    font_title = Font(name=font_family, size=15, bold=True, color="1F497D")
    font_subtitle = Font(name=font_family, size=10, italic=True, color="595959")
    font_section = Font(name=font_family, size=11, bold=True, color="1F497D")
    
    font_kpi_label = Font(name=font_family, size=9, bold=False, color="595959")
    font_kpi_val = Font(name=font_family, size=11, bold=True, color="000000")
    
    font_tbl_header = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    font_data = Font(name=font_family, size=9, color="000000")
    font_data_bold = Font(name=font_family, size=9, bold=True, color="000000")
    
    fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")
    fill_kpi_bg = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    fill_status_closed = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_status_open = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(bottom=Side(border_style="medium", color="1F497D"))
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # Plain numeric format strings (strictly no rupee symbol)
    FMT_CURRENCY = "#,##0.00"
    FMT_QTY = "#,##0"
    FMT_PCT = "0.00%"
    FMT_DATE = "DD-MMM-YYYY"

    # ---------------------------------------------------------
    # Title Section
    # ---------------------------------------------------------
    ws["B2"] = "WEEKLY BREAKOUT SWING STRATEGY - BACKTEST REPORT"
    ws["B2"].font = font_title
    ws["B3"] = f"Stock: {result.symbol} | Strategy: 5-Tranche 5% Profit Target | Capital: 100,000 | Max Tranches: 5"
    ws["B3"].font = font_subtitle

    # ---------------------------------------------------------
    # Top Executive Summary Dashboard (Cards Layout)
    # ---------------------------------------------------------
    ws["B5"] = "PORTFOLIO & CAPITAL SUMMARY"
    ws["B5"].font = font_section

    if config.enable_magic_profit:
        kpi_items = [
            ("Total Capital", result.initial_capital, FMT_CURRENCY),
            ("Total Realized Profit", result.total_realized_profit, FMT_CURRENCY),
            ("Trade Realized Profit", result.realized_trade_profit, FMT_CURRENCY),
            ("Magic Profit (+500)", result.total_magic_profit, FMT_CURRENCY),
            ("Unrealized P&L", result.unrealized_pnl, FMT_CURRENCY),
            ("Total Portfolio Return", result.total_portfolio_return_pct / 100.0, FMT_PCT),
            ("XIRR %", result.deployed_xirr_pct / 100.0, FMT_PCT),
            ("Max Drawdown", result.max_drawdown_pct / 100.0, FMT_PCT),
            ("Max Tranches Deployed", result.max_active_tranches, FMT_QTY),
            ("Total Trades Taken", result.total_trades, FMT_QTY),
            ("Completed Trades", result.completed_trades, FMT_QTY),
            ("Open Trades", result.open_trades, FMT_QTY),
            ("Magic Events Count", len(result.magic_events), FMT_QTY),
            ("Avg Holding Days", result.avg_holding_days, "0.0"),
        ]
        card_rows_count = 4
    else:
        kpi_items = [
            ("Total Capital", result.initial_capital, FMT_CURRENCY),
            ("Total Realized Profit", result.total_realized_profit, FMT_CURRENCY),
            ("Unrealized P&L", result.unrealized_pnl, FMT_CURRENCY),
            ("Total Portfolio Return", result.total_portfolio_return_pct / 100.0, FMT_PCT),
            ("XIRR %", result.deployed_xirr_pct / 100.0, FMT_PCT),
            ("Max Drawdown", result.max_drawdown_pct / 100.0, FMT_PCT),
            ("Capital Per Tranche", config.tranche_capital, FMT_CURRENCY),
            ("Max Tranches Deployed", result.max_active_tranches, FMT_QTY),
            ("Total Trades Taken", result.total_trades, FMT_QTY),
            ("Completed Trades", result.completed_trades, FMT_QTY),
            ("Open Trades", result.open_trades, FMT_QTY),
            ("Avg Holding Days", result.avg_holding_days, "0.0"),
            ("Open Invested Capital", result.open_positions_invested, FMT_CURRENCY),
            ("Max Active Slots", result.max_active_tranches, FMT_QTY),
        ]
        card_rows_count = 4

    # Render KPI cards in a 4-column block starting at B6
    card_cols = [("B", "C"), ("E", "F"), ("H", "I"), ("K", "L")]
    for idx, (label, val, fmt) in enumerate(kpi_items):
        row_offset = idx // 4
        col_pair = card_cols[idx % 4]
        c_label, c_val = col_pair[0], col_pair[1]
        r = 6 + (row_offset * 2)

        # Label cell
        cell_lbl = ws[f"{c_label}{r}"]
        cell_lbl.value = label
        cell_lbl.font = font_kpi_label
        cell_lbl.alignment = align_left
        cell_lbl.fill = fill_kpi_bg
        cell_lbl.border = thin_border

        # Value cell
        cell_val = ws[f"{c_val}{r}"]
        cell_val.value = val
        cell_val.font = font_kpi_val
        cell_val.alignment = align_right
        cell_val.fill = fill_kpi_bg
        cell_val.border = thin_border
        cell_val.number_format = fmt

    # ---------------------------------------------------------
    # Detailed Trade Transactions Table
    # ---------------------------------------------------------
    tbl_start_row = 6 + (card_rows_count * 2) + 2
    ws[f"B{tbl_start_row - 1}"] = "DETAILED TRADE TRANSACTIONS LOG"
    ws[f"B{tbl_start_row - 1}"].font = font_section

    if config.enable_magic_profit:
        headers = [
            ("Trade #", align_center),
            ("Tranche Slot", align_center),
            ("Symbol", align_left),
            ("Signal Date", align_center),
            ("Buy Date", align_center),
            ("Buy Price", align_right),
            ("Shares", align_right),
            ("Invested Amount", align_right),
            ("Target Price (5%)", align_right),
            ("Sell Date", align_center),
            ("Sell Price", align_right),
            ("Trade Realized P&L", align_right),
            ("Magic Profit", align_right),
            ("Total Realized P&L", align_right),
            ("Magic Trigger Dates", align_left),
            ("Return %", align_right),
            ("Holding Days", align_center),
            ("Status", align_center),
        ]
        status_col_idx = 2 + len(headers) - 1
    else:
        headers = [
            ("Trade #", align_center),
            ("Tranche Slot", align_center),
            ("Symbol", align_left),
            ("Signal Date", align_center),
            ("Buy Date", align_center),
            ("Buy Price", align_right),
            ("Shares", align_right),
            ("Invested Amount", align_right),
            ("Target Price (5%)", align_right),
            ("Sell Date", align_center),
            ("Sell Price", align_right),
            ("Realized P&L", align_right),
            ("Return %", align_right),
            ("Holding Days", align_center),
            ("Status", align_center),
        ]
        status_col_idx = 2 + len(headers) - 1

    # Write Table Headers
    for c_idx, (h_title, h_align) in enumerate(headers, start=2): # Start at col B (2)
        cell = ws.cell(row=tbl_start_row, column=c_idx, value=h_title)
        cell.font = font_tbl_header
        cell.fill = fill_header
        cell.alignment = h_align
        cell.border = thin_border

    # Write Trade Rows
    curr_row = tbl_start_row + 1
    for t in result.trades:
        is_zebra = (curr_row % 2 == 0)
        row_fill = fill_zebra if is_zebra else None

        signal_str = t.signal_date.strftime("%d-%b-%Y") if t.signal_date is not None else "-"
        buy_str = t.buy_date.strftime("%d-%b-%Y") if t.buy_date is not None else "-"
        sell_str = t.sell_date.strftime("%d-%b-%Y") if t.sell_date is not None else "-"
        ret_val = (t.return_pct / 100.0) if t.return_pct is not None else None

        if config.enable_magic_profit:
            total_tranche_pnl = (
                round((t.realized_pnl or 0.0) + t.magic_profit_total, 2)
                if t.realized_pnl is not None
                else (round(t.magic_profit_total, 2) if t.magic_profit_total > 0 else None)
            )
            magic_dates_str = ", ".join(d.strftime("%d-%b-%Y") for d in t.magic_dates) if t.magic_dates else "-"
            
            row_data = [
                (t.trade_id, FMT_QTY, align_center),
                (t.slot_id, FMT_QTY, align_center),
                (t.symbol, "@", align_left),
                (signal_str, FMT_DATE, align_center),
                (buy_str, FMT_DATE, align_center),
                (t.buy_price, FMT_CURRENCY, align_right),
                (t.shares, FMT_QTY, align_right),
                (t.invested_amt, FMT_CURRENCY, align_right),
                (t.target_price, FMT_CURRENCY, align_right),
                (sell_str, FMT_DATE, align_center),
                (t.sell_price if t.sell_price is not None else "-", FMT_CURRENCY if t.sell_price else "@", align_right),
                (t.realized_pnl if t.realized_pnl is not None else "-", FMT_CURRENCY if t.realized_pnl else "@", align_right),
                (t.magic_profit_total, FMT_CURRENCY, align_right),
                (total_tranche_pnl if total_tranche_pnl is not None else "-", FMT_CURRENCY if total_tranche_pnl is not None else "@", align_right),
                (magic_dates_str, "@", align_left),
                (ret_val if ret_val is not None else "-", FMT_PCT if ret_val else "@", align_right),
                (t.holding_days if t.holding_days is not None else "-", FMT_QTY if t.holding_days else "@", align_center),
                (t.status, "@", align_center),
            ]
        else:
            row_data = [
                (t.trade_id, FMT_QTY, align_center),
                (t.slot_id, FMT_QTY, align_center),
                (t.symbol, "@", align_left),
                (signal_str, FMT_DATE, align_center),
                (buy_str, FMT_DATE, align_center),
                (t.buy_price, FMT_CURRENCY, align_right),
                (t.shares, FMT_QTY, align_right),
                (t.invested_amt, FMT_CURRENCY, align_right),
                (t.target_price, FMT_CURRENCY, align_right),
                (sell_str, FMT_DATE, align_center),
                (t.sell_price if t.sell_price is not None else "-", FMT_CURRENCY if t.sell_price else "@", align_right),
                (t.realized_pnl if t.realized_pnl is not None else "-", FMT_CURRENCY if t.realized_pnl else "@", align_right),
                (ret_val if ret_val is not None else "-", FMT_PCT if ret_val else "@", align_right),
                (t.holding_days if t.holding_days is not None else "-", FMT_QTY if t.holding_days else "@", align_center),
                (t.status, "@", align_center),
            ]

        for c_idx, (val, fmt, cell_align) in enumerate(row_data, start=2):
            cell = ws.cell(row=curr_row, column=c_idx, value=val)
            cell.font = font_data
            cell.alignment = cell_align
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
            if fmt != "@" and isinstance(val, (int, float)):
                cell.number_format = fmt

            # Specific badge coloring for Status column
            if c_idx == status_col_idx:
                if val == "CLOSED":
                    cell.fill = fill_status_closed
                    cell.font = font_data_bold
                elif val == "OPEN":
                    cell.fill = fill_status_open
                    cell.font = font_data_bold

        curr_row += 1

    # Freeze pane so headers stay fixed during scrolling
    ws.freeze_panes = f"B{tbl_start_row + 1}"

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col[0].column == 1:
            ws.column_dimensions[col_letter].width = 3
            continue
        max_len = 0
        for cell in col:
            if cell.row < tbl_start_row - 2:
                continue  # skip title/kpis for table width calc
            if cell.value:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    # ---------------------------------------------------------
    # Dedicated Magic Profit Journal Worksheet
    # ---------------------------------------------------------
    if config.enable_magic_profit and result.magic_events:
        ws_magic = wb.create_sheet(title="Magic Profit Log")
        ws_magic.views.sheetView[0].showGridLines = True

        ws_magic["B2"] = "MAGIC PROFIT EVENT JOURNAL (+500 ON 20-DAY LOW BREAKDOWN)"
        ws_magic["B2"].font = font_title
        ws_magic["B3"] = f"Stock: {result.symbol} | Trigger: Close < 20DL (upto yesterday) while Tranches in Loss | Cooldown: {config.magic_profit_cooldown_days} Days | Credit: 500 once for all open tranches"
        ws_magic["B3"].font = font_subtitle

        m_tbl_start = 5
        m_headers = [
            ("Event #", align_center),
            ("Event Date", align_center),
            ("Credited Trade #", align_center),
            ("Tranche Slot", align_center),
            ("Tranche Buy Date", align_center),
            ("Tranche Buy Price", align_right),
            ("20DL (upto yesterday)", align_right),
            ("Breakdown Close Price", align_right),
            ("Loss % at Trigger", align_right),
            ("Open Tranches in Loss", align_left),
            ("Profit Credited", align_right),
            ("Cumulative Magic Profit", align_right),
        ]

        for c_idx, (h_title, h_align) in enumerate(m_headers, start=2):
            cell = ws_magic.cell(row=m_tbl_start, column=c_idx, value=h_title)
            cell.font = font_tbl_header
            cell.fill = fill_header
            cell.alignment = h_align
            cell.border = thin_border

        m_curr_row = m_tbl_start + 1
        cum_magic = 0.0
        for ev in result.magic_events:
            cum_magic += ev.profit_credit
            is_zebra = (m_curr_row % 2 == 0)
            row_fill = fill_zebra if is_zebra else None

            ev_date_str = ev.date.strftime("%d-%b-%Y") if ev.date is not None else "-"
            b_date_str = ev.buy_date.strftime("%d-%b-%Y") if ev.buy_date is not None else "-"

            m_row_data = [
                (ev.event_id, FMT_QTY, align_center),
                (ev_date_str, FMT_DATE, align_center),
                (ev.trade_id, FMT_QTY, align_center),
                (ev.slot_id, FMT_QTY, align_center),
                (b_date_str, FMT_DATE, align_center),
                (ev.buy_price, FMT_CURRENCY, align_right),
                (ev.rolling_20d_low, FMT_CURRENCY, align_right),
                (ev.close_price, FMT_CURRENCY, align_right),
                (ev.loss_pct / 100.0, FMT_PCT, align_right),
                (ev.open_tranches_in_loss or f"Slot {ev.slot_id}", "@", align_left),
                (ev.profit_credit, FMT_CURRENCY, align_right),
                (cum_magic, FMT_CURRENCY, align_right),
            ]

            for c_idx, (val, fmt, cell_align) in enumerate(m_row_data, start=2):
                cell = ws_magic.cell(row=m_curr_row, column=c_idx, value=val)
                cell.font = font_data
                cell.alignment = cell_align
                cell.border = thin_border
                if row_fill:
                    cell.fill = row_fill
                if fmt != "@" and isinstance(val, (int, float)):
                    cell.number_format = fmt

            m_curr_row += 1

        ws_magic.freeze_panes = f"B{m_tbl_start + 1}"
        for col in ws_magic.columns:
            col_letter = get_column_letter(col[0].column)
            if col[0].column == 1:
                ws_magic.column_dimensions[col_letter].width = 3
                continue
            max_len = 0
            for cell in col:
                if cell.row < m_tbl_start:
                    continue
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws_magic.column_dimensions[col_letter].width = max(max_len + 4, 13)

    try:
        wb.save(output_filepath)
        return output_filepath
    except PermissionError:
        base, ext = os.path.splitext(output_filepath)
        fallback_path = f"{base}_new{ext}"
        wb.save(fallback_path)
        print(f"[WARNING] {output_filepath} is currently open in Excel. Saved updated report to {fallback_path} instead.")
        return fallback_path
